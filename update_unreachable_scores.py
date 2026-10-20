"""
update_unreachable_scores.py
อัปเดตคะแนนเว็บไซต์ที่เข้าไม่ได้จริง (ตาย / 404 / พัง / โดเมนหมดอายุ)
ให้ได้ 0 คะแนน (0%) เกรด "ต้องปรับปรุง" ใน standards_score_results.xlsx ทันที
"""
import sys
import openpyxl
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

BASE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE_DIR))

from score_scan import update_summary_dashboard

EXCEL_PATH = BASE_DIR / "standards_score_results.xlsx"
REACH_PATH = BASE_DIR / "reachability_report.xlsx"

def main():
    if not EXCEL_PATH.exists() or not REACH_PATH.exists():
        print("[!] ไม่พบไฟล์ excel ที่จำเป็น")
        return

    print("=" * 70)
    print("  อัปเดตคะแนนเว็บที่เข้าไม่ได้จริงให้เป็น 0% (เกรด ต้องปรับปรุง)")
    print("=" * 70)

    # 1. โหลดรายชื่อเว็บที่เข้าไม่ได้จาก reachability_report.xlsx
    wb_reach = openpyxl.load_workbook(REACH_PATH)
    ws_unreach = wb_reach["เข้าไม่ได้ (Unreachable)"]
    
    unreachable_map = {}
    for r in range(2, ws_unreach.max_row + 1):
        url = ws_unreach.cell(row=r, column=2).value
        err_type = ws_unreach.cell(row=r, column=5).value
        duration = ws_unreach.cell(row=r, column=7).value or 1.5
        if url:
            clean_u = str(url).strip().rstrip("/").lower()
            unreachable_map[clean_u] = {
                "original": str(url).strip(),
                "err_type": err_type,
                "duration": duration,
            }

    # เพิ่ม fpo.go.th/S-I/ เข้า unreachable ถ้ายังไม่มี
    unreachable_map["http://www.fpo.go.th/s-i"] = {
        "original": "http://www.fpo.go.th/S-I/",
        "err_type": "หน้าเว็บไม่มีอยู่จริง (HTTP 404)",
        "duration": 1.2
    }
    unreachable_map["https://www.fpo.go.th/s-i"] = {
        "original": "https://www.fpo.go.th/S-I/",
        "err_type": "หน้าเว็บไม่มีอยู่จริง (HTTP 404)",
        "duration": 1.2
    }

    # fpo.go.th หน้าหลักเป็นเว็บจริง ต้องไม่ถูกให้ 0
    unreachable_map.pop("http://www.fpo.go.th", None)
    unreachable_map.pop("https://www.fpo.go.th", None)

    print(f"[*] พบเว็บที่เข้าไม่ได้จริงทั้งหมด {len(unreachable_map)} เว็บไซต์")

    # 2. โหลด standards_score_results.xlsx
    wb_std = openpyxl.load_workbook(EXCEL_PATH)
    ws_std = wb_std["Standards Score Report"]

    updated_count = 0
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    red_fill = PatternFill(start_color="F8D7DA", fill_type="solid")
    red_font = Font(name="Segoe UI", size=10, bold=True, color="842029")
    status_fail_fill = PatternFill(start_color="FCE4D6", fill_type="solid")
    status_fail_font = Font(name="Segoe UI", size=10, bold=True, color="C00000")

    for r in range(2, ws_std.max_row + 1):
        raw_url = ws_std.cell(row=r, column=2).value
        if not raw_url:
            continue
        clean_u = str(raw_url).strip().rstrip("/").lower()

        if clean_u in unreachable_map:
            info = unreachable_map[clean_u]
            updated_count += 1

            # สูตรคะแนน
            ws_std.cell(row=r, column=4).value = f"=F{r}/71"
            ws_std.cell(row=r, column=5).value = f'=IF(D{r}>=0.9,"ดีเยี่ยม",IF(D{r}>=0.7,"ดี",IF(D{r}>=0.5,"พอใช้","ต้องปรับปรุง")))'
            ws_std.cell(row=r, column=6).value = f"=J{r}+N{r}+R{r}+V{r}"
            ws_std.cell(row=r, column=7).value = f"=K{r}+O{r}+S{r}+W{r}"
            ws_std.cell(row=r, column=8).value = f"=L{r}+P{r}+T{r}+X{r}"

            # WCAG (0/37)
            ws_std.cell(row=r, column=9).value = f"=J{r}/37"
            ws_std.cell(row=r, column=10).value = 0
            ws_std.cell(row=r, column=11).value = 37
            ws_std.cell(row=r, column=12).value = 0

            # CWV (0/9)
            ws_std.cell(row=r, column=13).value = f"=N{r}/9"
            ws_std.cell(row=r, column=14).value = 0
            ws_std.cell(row=r, column=15).value = 9
            ws_std.cell(row=r, column=16).value = 0

            # NCSA (0/11)
            ws_std.cell(row=r, column=17).value = f"=R{r}/11"
            ws_std.cell(row=r, column=18).value = 0
            ws_std.cell(row=r, column=19).value = 11
            ws_std.cell(row=r, column=20).value = 0

            # OWASP (0/14)
            ws_std.cell(row=r, column=21).value = f"=V{r}/14"
            ws_std.cell(row=r, column=22).value = 0
            ws_std.cell(row=r, column=23).value = 14
            ws_std.cell(row=r, column=24).value = 0

            # Duration, Status, Timestamp
            ws_std.cell(row=r, column=25).value = round(float(info["duration"]), 2)
            ws_std.cell(row=r, column=26).value = f"Failed ({info['err_type']})"
            ws_std.cell(row=r, column=27).value = now_str

            # ปรับสีให้เป็นสีแดง (ต้องปรับปรุง)
            c4 = ws_std.cell(row=r, column=4)
            c4.fill = red_fill
            c4.font = red_font
            
            c5 = ws_std.cell(row=r, column=5)
            c5.fill = red_fill
            c5.font = red_font

            c26 = ws_std.cell(row=r, column=26)
            c26.fill = status_fail_fill
            c26.font = status_fail_font

            print(f"  [{updated_count:3d}] ปรับเป็น 0%: {raw_url:45s} | {info['err_type']}")

    print(f"\n[*] อัปเดตครบทั้งหมด {updated_count} เว็บไซต์")
    print("[*] กำลังคำนวณชีต Summary สถิติภาพรวม ใหม่...")
    update_summary_dashboard(wb_std)

    wb_std.save(EXCEL_PATH)
    print(f"[✅] บันทึกไฟล์สำเร็จ → {EXCEL_PATH.name}")
    print("=" * 70)

if __name__ == "__main__":
    main()
