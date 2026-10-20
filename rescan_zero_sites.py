"""
rescan_zero_sites.py
Identifies all websites in F:\\Potato\\standards_score_results.xlsx that had
WCAG 0% (due to previous redirect/navigation issues), re-scans them with the
updated redirect-resilient scanner, and updates the rows and dashboard in place.
"""
import os
import sys
import time
import json
import tempfile
import urllib.request
from datetime import datetime
from pathlib import Path

# Redirect temp directory to Drive F: to keep Drive C: clean
F_TEMP = Path(r"F:\Potato\temp")
F_TEMP.mkdir(parents=True, exist_ok=True)
os.environ["TEMP"] = str(F_TEMP)
os.environ["TMP"] = str(F_TEMP)
tempfile.tempdir = str(F_TEMP)

BASE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE_DIR))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from score_scan import (
    get_grade,
    update_summary_dashboard,
    scan_via_backend_api,
)

EXCEL_PATH = BASE_DIR / "standards_score_results.xlsx"


def parse_report_for_row_meta(report: dict, url: str, category: str, row_num: int, duration: float) -> dict:
    summary = report.get("summary", {}) if isinstance(report, dict) else {}
    total_items = summary.get("total", 71)
    passed_items = summary.get("passed", 0)
    failed_items = summary.get("failed", 0)
    warning_items = summary.get("warning", 0)

    overall_score = round((passed_items / total_items * 100)) if total_items > 0 else 0
    grade_label, _, _ = get_grade(overall_score)

    standards = report.get("standards", []) if isinstance(report, dict) else []
    std_map = {s.get("id"): s for s in standards if isinstance(s, dict)}

    wcag = std_map.get("wcag", {"passed": 0, "failed": 0, "warning": 0, "total": 37})
    cwv = std_map.get("cwv", {"passed": 0, "failed": 0, "warning": 0, "total": 9})
    ncsa = std_map.get("ncsa", {"passed": 0, "failed": 0, "warning": 0, "total": 11})
    owasp = std_map.get("owasp", {"passed": 0, "failed": 0, "warning": 0, "total": 14})

    return {
        "row_num": row_num,
        "url": url,
        "category": category,
        "overall_score": overall_score,
        "grade": grade_label,
        "passed": passed_items,
        "failed": failed_items,
        "warning": warning_items,
        "total": total_items,
        # WCAG
        "wcag_passed": wcag.get("passed", 0),
        "wcag_failed": wcag.get("failed", 0),
        "wcag_warning": wcag.get("warning", 0),
        # CWV
        "cwv_passed": cwv.get("passed", 0),
        "cwv_failed": cwv.get("failed", 0),
        "cwv_warning": cwv.get("warning", 0),
        # NCSA
        "ncsa_passed": ncsa.get("passed", 0),
        "ncsa_failed": ncsa.get("failed", 0),
        "ncsa_warning": ncsa.get("warning", 0),
        # OWASP
        "owasp_passed": owasp.get("passed", 0),
        "owasp_failed": owasp.get("failed", 0),
        "owasp_warning": owasp.get("warning", 0),
        # Meta
        "scan_duration": duration,
        "status": "Success" if (wcag.get("passed", 0) > 0 or wcag.get("failed", 0) > 0) else "Unreachable",
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def update_excel_row(ws, r: int, row_meta: dict):
    """Update row cells in place, preserving dynamic formulas and applying styling."""
    bold_font = Font(name="Segoe UI", size=10, bold=True)

    # WCAG (cols 10, 11, 12)
    ws.cell(row=r, column=10, value=row_meta["wcag_passed"])
    ws.cell(row=r, column=11, value=row_meta["wcag_failed"])
    ws.cell(row=r, column=12, value=row_meta["wcag_warning"])

    # CWV (cols 14, 15, 16)
    ws.cell(row=r, column=14, value=row_meta["cwv_passed"])
    ws.cell(row=r, column=15, value=row_meta["cwv_failed"])
    ws.cell(row=r, column=16, value=row_meta["cwv_warning"])

    # NCSA (cols 18, 19, 20)
    ws.cell(row=r, column=18, value=row_meta["ncsa_passed"])
    ws.cell(row=r, column=19, value=row_meta["ncsa_failed"])
    ws.cell(row=r, column=20, value=row_meta["ncsa_warning"])

    # OWASP (cols 22, 23, 24)
    ws.cell(row=r, column=22, value=row_meta["owasp_passed"])
    ws.cell(row=r, column=23, value=row_meta["owasp_failed"])
    ws.cell(row=r, column=24, value=row_meta["owasp_warning"])

    # Meta
    ws.cell(row=r, column=25, value=row_meta["scan_duration"])
    ws.cell(row=r, column=26, value=row_meta["status"])
    ws.cell(row=r, column=27, value=row_meta["timestamp"])

    # Format percentages as bold
    for c_idx in [4, 9, 13, 17, 21]:
        ws.cell(row=r, column=c_idx).font = bold_font
        ws.cell(row=r, column=c_idx).number_format = '0%'

    # Highlight Overall Score & Grade
    pct_num = row_meta.get("overall_score", 0)
    _, bg, fg = get_grade(pct_num)
    ws.cell(row=r, column=4).fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    ws.cell(row=r, column=4).font = Font(name="Segoe UI", size=10, bold=True, color=fg)

    grade_str = row_meta.get("grade", "")
    grade_cell = ws.cell(row=r, column=5)
    if grade_str == "ดีเยี่ยม":
        grade_cell.fill = PatternFill(start_color="D4EDDA", fill_type="solid")
        grade_cell.font = Font(name="Segoe UI", size=10, bold=True, color="155724")
    elif grade_str == "ดี":
        grade_cell.fill = PatternFill(start_color="D1E7DD", fill_type="solid")
        grade_cell.font = Font(name="Segoe UI", size=10, bold=True, color="0F5132")
    elif grade_str == "พอใช้":
        grade_cell.fill = PatternFill(start_color="FFF3CD", fill_type="solid")
        grade_cell.font = Font(name="Segoe UI", size=10, bold=True, color="856404")
    else:
        grade_cell.fill = PatternFill(start_color="F8D7DA", fill_type="solid")
        grade_cell.font = Font(name="Segoe UI", size=10, bold=True, color="842029")

    # Status Badge
    status_cell = ws.cell(row=r, column=26)
    val = str(status_cell.value)
    if "Success" in val:
        status_cell.fill = PatternFill(start_color="E2EFDA", fill_type="solid")
        status_cell.font = Font(name="Segoe UI", size=10, color="276A3C", bold=True)
    elif "Partial" in val:
        status_cell.fill = PatternFill(start_color="FFF2CC", fill_type="solid")
        status_cell.font = Font(name="Segoe UI", size=10, color="8A6D3B", bold=True)
    else:
        status_cell.fill = PatternFill(start_color="FCE4D6", fill_type="solid")
        status_cell.font = Font(name="Segoe UI", size=10, color="C00000", bold=True)


def main():
    if not EXCEL_PATH.exists():
        print(f"[!] File not found: {EXCEL_PATH}")
        sys.exit(1)

    print("=" * 68)
    print("  Re-Scanning Websites with 0% WCAG (Redirect-Resilient Mode)")
    print("=" * 68)
    print(f"[*] Target Excel File: {EXCEL_PATH}")

    try:
        wb = openpyxl.load_workbook(EXCEL_PATH)
    except PermissionError:
        print(f"[!] กรุณาปิดไฟล์ {EXCEL_PATH.name} ในโปรแกรม Excel ก่อนเริ่มรัน!")
        sys.exit(1)

    ws = wb["Standards Score Report"]
    targets = []

    # Identify rows where WCAG had 0 passes and 0 fails (skipped due to redirect)
    for r in range(2, ws.max_row + 1):
        url = ws.cell(row=r, column=2).value
        category = ws.cell(row=r, column=3).value or "ทั่วไป"
        wcag_pass = ws.cell(row=r, column=10).value or 0
        wcag_fail = ws.cell(row=r, column=11).value or 0
        if wcag_pass == 0 and wcag_fail == 0:
            targets.append((r, str(url).strip(), str(category)))

    print(f"[*] Found {len(targets)} websites to verify and re-scan.")
    if not targets:
        print("[+] No 0% WCAG rows found! All rows are already verified.")
        return

    improved_count = 0
    offline_count = 0

    for idx, (r, url, cat) in enumerate(targets, start=1):
        print(f"\n[{idx}/{len(targets)}] Re-evaluating: {url}")
        t0 = time.perf_counter()

        report, dur = scan_via_backend_api(url, timeout=120)
        elapsed = round(time.perf_counter() - t0, 2)

        if report and isinstance(report, dict):
            row_meta = parse_report_for_row_meta(
                report=report,
                url=url,
                category=cat,
                row_num=r - 1,
                duration=elapsed,
            )

            wcag_p = row_meta.get("wcag_passed", 0)
            overall_pct = row_meta.get("overall_score", 0)
            grade = row_meta.get("grade", "")

            if wcag_p > 0 or row_meta.get("wcag_failed", 0) > 0:
                improved_count += 1
                row_meta["status"] = "Success"
                print(f"    ├─ Result: SUCCESS! WCAG: {wcag_p}/37 passed | Overall: {overall_pct}% ({grade}) | Duration: {elapsed}s")
            else:
                row_meta["status"] = "Unreachable"
                offline_count += 1
                print(f"    ├─ Result: Offline / No Content (WCAG: 0/37) | Overall: {overall_pct}%")

            update_excel_row(ws, r, row_meta)
        else:
            offline_count += 1
            ws.cell(row=r, column=26, value="Offline/DNS Error")
            ws.cell(row=r, column=26).fill = PatternFill(start_color="FCE4D6", fill_type="solid")
            ws.cell(row=r, column=26).font = Font(name="Segoe UI", size=10, color="C00000", bold=True)
            print(f"    └─ Result: Site Dead / Unreachable (Duration: {elapsed}s)")

        # Save checkpoint every 5 sites
        if idx % 5 == 0 or idx == len(targets):
            update_summary_dashboard(wb)
            try:
                wb.save(EXCEL_PATH)
                print(f"    [*] Checkpoint saved ({idx}/{len(targets)} completed)")
            except PermissionError:
                print(f"    [!] Warning: Excel is open, skipping autosave checkpoint...")

    update_summary_dashboard(wb)
    wb.save(EXCEL_PATH)
    print("\n" + "=" * 68)
    print(f"[OK] Verification & Re-scan Complete!")
    print(f"     Total re-checked : {len(targets)} websites")
    print(f"     Improved/Rescued : {improved_count} websites (now have full valid scores)")
    print(f"     Confirmed Offline: {offline_count} websites (confirmed dead/unreachable)")
    print(f"     File Updated     : {EXCEL_PATH}")
    print("=" * 68)


if __name__ == "__main__":
    main()
