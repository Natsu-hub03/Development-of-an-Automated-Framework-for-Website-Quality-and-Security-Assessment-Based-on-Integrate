"""
smart_audit_repair.py
Comprehensive Audit & Repair Script for all 600 sites in standards_score_results.xlsx.
Phase 1: Parallel Reachability Audit of all 600 sites.
Phase 2: Instant zero-out for dead/closed/403/crash sites that currently leak points.
Phase 3: Targeted re-scan for reachable sites that were falsely zeroed or have missing WCAG.
"""
import os
import sys
import time
import shutil
from datetime import datetime
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor

# Ensure UTF-8 output on Windows consoles
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

BASE_DIR = Path(__file__).resolve().parent
BACKEND_DIR = BASE_DIR / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import openpyxl
from openpyxl.styles import Font

from services.reachability import check_url
from score_scan import evaluate_website_standards

EXCEL_PATH = BASE_DIR / "standards_score_results.xlsx"


def update_row_zero(ws, r: int, url: str, cat: str, reason: str):
    """Zeroes out an unreachable/dead website row in Excel."""
    cell_b = ws.cell(r, 2)
    cell_b.value = url
    cell_b.hyperlink = url if url.startswith("http") else f"http://{url}"
    cell_b.font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")

    ws.cell(r, 3).value = cat

    # Formulas
    ws.cell(r, 4).value = f"=F{r}/71"
    ws.cell(r, 5).value = f'=IF(D{r}>=0.9,"ดีเยี่ยม",IF(D{r}>=0.7,"ดี",IF(D{r}>=0.5,"พอใช้","ต้องปรับปรุง")))'
    ws.cell(r, 6).value = f"=J{r}+N{r}+R{r}+V{r}"
    ws.cell(r, 7).value = f"=K{r}+O{r}+S{r}+W{r}"
    ws.cell(r, 8).value = f"=L{r}+P{r}+T{r}+X{r}"

    # WCAG (0/37)
    ws.cell(r, 9).value = f"=J{r}/37"
    ws.cell(r, 10).value = 0
    ws.cell(r, 11).value = 37
    ws.cell(r, 12).value = 0

    # CWV (0/9)
    ws.cell(r, 13).value = f"=N{r}/9"
    ws.cell(r, 14).value = 0
    ws.cell(r, 15).value = 9
    ws.cell(r, 16).value = 0

    # NCSA (0/11)
    ws.cell(r, 17).value = f"=R{r}/11"
    ws.cell(r, 18).value = 0
    ws.cell(r, 19).value = 11
    ws.cell(r, 20).value = 0

    # OWASP (0/14)
    ws.cell(r, 21).value = f"=V{r}/14"
    ws.cell(r, 22).value = 0
    ws.cell(r, 23).value = 14
    ws.cell(r, 24).value = 0

    # Meta
    ws.cell(r, 25).value = 0.0
    ws.cell(r, 26).value = f"Dead/Closed ({reason})"
    ws.cell(r, 27).value = datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def update_row_full(ws, r: int, url: str, cat: str, res: dict):
    """Updates an alive website row with full evaluated scores."""
    cell_b = ws.cell(r, 2)
    cell_b.value = url
    cell_b.hyperlink = url if url.startswith("http") else f"http://{url}"
    cell_b.font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")

    ws.cell(r, 3).value = cat

    # Formulas
    ws.cell(r, 4).value = f"=F{r}/71"
    ws.cell(r, 5).value = f'=IF(D{r}>=0.9,"ดีเยี่ยม",IF(D{r}>=0.7,"ดี",IF(D{r}>=0.5,"พอใช้","ต้องปรับปรุง")))'
    ws.cell(r, 6).value = f"=J{r}+N{r}+R{r}+V{r}"
    ws.cell(r, 7).value = f"=K{r}+O{r}+S{r}+W{r}"
    ws.cell(r, 8).value = f"=L{r}+P{r}+T{r}+X{r}"

    # WCAG
    ws.cell(r, 9).value = f"=J{r}/37"
    ws.cell(r, 10).value = res["wcag_passed"]
    ws.cell(r, 11).value = res["wcag_failed"]
    ws.cell(r, 12).value = res["wcag_warning"]

    # CWV
    ws.cell(r, 13).value = f"=N{r}/9"
    ws.cell(r, 14).value = res["cwv_passed"]
    ws.cell(r, 15).value = res["cwv_failed"]
    ws.cell(r, 16).value = res["cwv_warning"]

    # NCSA
    ws.cell(r, 17).value = f"=R{r}/11"
    ws.cell(r, 18).value = res["ncsa_passed"]
    ws.cell(r, 19).value = res["ncsa_failed"]
    ws.cell(r, 20).value = res["ncsa_warning"]

    # OWASP
    ws.cell(r, 21).value = f"=V{r}/14"
    ws.cell(r, 22).value = res["owasp_passed"]
    ws.cell(r, 23).value = res["owasp_failed"]
    ws.cell(r, 24).value = res["owasp_warning"]

    # Meta
    ws.cell(r, 25).value = res["scan_duration"]
    ws.cell(r, 26).value = res["status"]
    ws.cell(r, 27).value = datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def main():
    print("=" * 68)
    print("  600-Website Smart Audit & Precision Repair")
    print("=" * 68)

    if not EXCEL_PATH.exists():
        print(f"[!] File not found: {EXCEL_PATH}")
        sys.exit(1)

    backup_name = f"standards_score_results_backup_before_audit_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    backup_path = BASE_DIR / backup_name
    shutil.copy2(EXCEL_PATH, backup_path)
    print(f"[*] สำรองไฟล์เดิมไว้ที่ -> {backup_name}")

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["Standards Score Report"]

    total_rows = ws.max_row - 1
    print(f"[*] จำนวนเว็บไซต์ในระบบ : {total_rows} เว็บไซต์")

    # Read existing rows
    items = []
    for r in range(2, ws.max_row + 1):
        num = ws.cell(r, 1).value
        url = str(ws.cell(r, 2).value or "").strip()
        cat = str(ws.cell(r, 3).value or "").strip()
        wcag_p = ws.cell(r, 10).value or 0
        wcag_f = ws.cell(r, 11).value or 0
        cwv_p = ws.cell(r, 14).value or 0
        ncsa_p = ws.cell(r, 18).value or 0
        owasp_p = ws.cell(r, 22).value or 0
        tot_p = wcag_p + cwv_p + ncsa_p + owasp_p
        status = str(ws.cell(r, 26).value or "").strip()
        items.append({
            "row": r,
            "no": num,
            "url": url,
            "cat": cat,
            "wcag_p": wcag_p,
            "wcag_f": wcag_f,
            "cwv_p": cwv_p,
            "tot_p": tot_p,
            "status": status,
        })

    # ── Phase 1: Parallel Reachability Audit (10 threads) ─────────────────
    print("\n[Phase 1/3] เริ่มตรวจสอบการเข้าถึงทั้ง 600 เว็บไซต์แบบขนาน (10 threads)...")
    t0 = time.perf_counter()

    def audit_url(item):
        reach = check_url(item["url"])
        return item, reach

    audit_results = []
    with ThreadPoolExecutor(max_workers=10) as executor:
        for res in executor.map(audit_url, items):
            audit_results.append(res)

    audit_dur = round(time.perf_counter() - t0, 2)
    print(f"[*] ตรวจสอบสถานะครบ 600 เว็บไซต์แล้ว ใช้เวลาเพียง: {audit_dur}s")

    # Classify items
    dead_to_zero = []
    alive_to_rescan = []
    already_correct = 0

    for item, reach in audit_results:
        is_reachable = reach.get("reachable", False)
        err = reach.get("error_type") or reach.get("error") or "เข้าถึงไม่ได้"

        if not is_reachable:
            # Case 1: Dead site, but currently has score > 0 -> Must be zeroed out
            if item["tot_p"] > 0:
                dead_to_zero.append((item, err))
            else:
                already_correct += 1
        else:
            # Case 2: Reachable site, but was falsely marked as 0% OR missing WCAG scan
            if item["tot_p"] == 0:
                alive_to_rescan.append((item, "Falsely zeroed (Site is alive)"))
            elif item["wcag_p"] == 0 and item["wcag_f"] == 0:
                alive_to_rescan.append((item, "Missing WCAG scan (0/0)"))
            elif "Partial" in item["status"]:
                alive_to_rescan.append((item, f"Incomplete status: {item['status']}"))
            else:
                already_correct += 1

    print(f"\n[*] สรุปผลการตรวจสอบ:")
    print(f"    ├─ เว็บไซต์ที่ถูกต้องสมบูรณ์อยู่แล้ว : {already_correct} เว็บ")
    print(f"    ├─ เว็บไซต์ตายจริงที่ต้องปรับเป็น 0    : {len(dead_to_zero)} เว็บ (เช่น share2trade.com, mobifan.com)")
    print(f"    └─ เว็บไซต์เปิดได้จริงที่ต้องสแกนใหม่ : {len(alive_to_rescan)} เว็บ (เช่น onec.go.th, openrice.com)")

    # ── Phase 2: Instant zero-out for dead sites ──────────────────────────
    print(f"\n[Phase 2/3] กำลังปรับคะแนนเป็น 0 สำหรับเว็บที่ตายจริง ({len(dead_to_zero)} เว็บ)...")
    for item, err in dead_to_zero:
        r = item["row"]
        update_row_zero(ws, r, item["url"], item["cat"], err)
        print(f"    [Zeroed] แถว {r} (No. {item['no']}): {item['url']} -> 0 คะแนนทุกมาตรฐาน ({err})")

    wb.save(EXCEL_PATH)
    print(f"[*] บันทึกเว็บที่ตายจริงเป็น 0 เรียบร้อยแล้ว!")

    # ── Phase 3: Targeted re-scan for alive sites needing scores ──────────
    if alive_to_rescan:
        print(f"\n[Phase 3/3] กำลังสแกนประเมินคะแนนใหม่สำหรับเว็บที่เปิดได้จริง ({len(alive_to_rescan)} เว็บ)...")
        timeouts = {
            "axe": 90,
            "lighthouse": 130,
            "headers": 45,
            "wappalyzer": 100,
        }

        for idx, (item, reason) in enumerate(alive_to_rescan, start=1):
            r = item["row"]
            url = item["url"]
            print(f"\n[{idx}/{len(alive_to_rescan)}] สแกนซ่อมแถว {r} (No. {item['no']}): {url}")
            print(f"    └─ สาเหตุ: {reason}")

            try:
                res = evaluate_website_standards(url, timeouts)
            except Exception as e:
                print(f"    [!] Error during scan: {e}")
                res = {
                    "overall_score": 0,
                    "grade": "ต้องปรับปรุง",
                    "passed": 0,
                    "failed": 71,
                    "warning": 0,
                    "wcag_passed": 0,
                    "wcag_failed": 37,
                    "wcag_warning": 0,
                    "cwv_passed": 0,
                    "cwv_failed": 9,
                    "cwv_warning": 0,
                    "ncsa_passed": 0,
                    "ncsa_failed": 11,
                    "ncsa_warning": 0,
                    "owasp_passed": 0,
                    "owasp_failed": 14,
                    "owasp_warning": 0,
                    "scan_duration": 0.0,
                    "status": f"Dead/Closed ({type(e).__name__})",
                }

            update_row_full(ws, r, url, item["cat"], res)
            wb.save(EXCEL_PATH)

            print(f"    ├─ ผลลัพธ์ใหม่       : คะแนน {res['overall_score']}% ({res['grade']}) | ผ่าน: {res['passed']}/71 ตก: {res['failed']}")
            print(f"    ├─ WCAG 2.1         : ผ่าน {res['wcag_passed']}/37, ตก {res['wcag_failed']}")
            print(f"    ├─ CWV & SEO        : ผ่าน {res['cwv_passed']}/9, ตก {res['cwv_failed']}")
            print(f"    ├─ สกมช. NCSA       : ผ่าน {res['ncsa_passed']}/11, ตก {res['ncsa_failed']}")
            print(f"    ├─ OWASP Headers    : ผ่าน {res['owasp_passed']}/14, ตก {res['owasp_failed']}")
            print(f"    ├─ สถานะ            : {res['status']}")
            print(f"    └─ Saved row {r} to Excel checkpoint!")

    wb.save(EXCEL_PATH)
    print("\n" + "=" * 68)
    print("Audit & Precision Repair Complete!")
    print(f"Updated Excel file saved at: {EXCEL_PATH}")
    print("=" * 68)


if __name__ == "__main__":
    main()
