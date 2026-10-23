"""
repair_user_targets.py
Targeted audit and rescan for the user's flagged sites in standards_score_results.xlsx:
- Rescan yellow-highlighted sites so WCAG gets its real compliance score.
- Rescan alive red-highlighted sites so they get their full standards scores instead of 0.
- Ensure parked domains (Atom, GoDaddy, HugeDomains), blank redirect loops, and dead sites are strictly 0% (เว็บปิด, สีแดง).
"""

import os
import sys
import json
import time
import subprocess
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
BASE_DIR = Path(__file__).resolve().parent
BACKEND_DIR = BASE_DIR / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from services.standards.report_builder import build_standards_report, build_unreachable_standards_report

SCANNER_DIR = BASE_DIR / "scanner"
EXCEL_PATH = BASE_DIR / "standards_score_results.xlsx"

# 1. Sites that are confirmed DEAD / PARKED / REDIRECT-LOOP -> Set to 0% เว็บปิด
DEAD_SITES = {
    "http://thaigaming.com": "เว็บปิด (หน้าขาว / ติด Redirect Loop)",
    "http://www.caraudionetwork.com": "เว็บปิด (หน้าประกาศขายโดเมน Atom.com)",
    "http://www.zuzaa.com": "เว็บปิด (เว็บเดิมปิดบริการ ดีดไป TikTok)",
    "http://thaitechno.net": "เว็บปิด (Server Error 500 / เซิร์ฟเวอร์ไม่ตอบสนอง)",
}

# 2. Sites to RESCAN fully
RESCAN_SITES = [
    "http://feedforfuture.co",
    "http://www.nmm.ac.th",
    "http://www.thailandcube.com",
    "http://www.condonewb.com",
    "http://www.thaiclinic.com",
    "http://www.koratdaily.com",
    "http://www.propertyhub.in.th",
    "http://www.moj.go.th",
    "http://www.bectero.com",
    "http://www.prsociety.net",
    "http://www.chibimonchronicle.net",
    "http://www.getcools.com",
]

def run_node(script: str, url: str, timeout: int = 40):
    cmd = ["node", str(SCANNER_DIR / script), url]
    env = os.environ.copy()
    chrome_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
    if os.path.exists(chrome_path):
        env["PUPPETEER_EXECUTABLE_PATH"] = chrome_path
        env["CHROME_PATH"] = chrome_path
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout, cwd=str(SCANNER_DIR), env=env)
        out = p.stdout.strip()
        if "{" in out and "}" in out:
            json_str = out[out.find("{"):out.rfind("}")+1]
            return json.loads(json_str)
        return {"error": p.stderr.strip() or "No JSON output"}
    except Exception as e:
        return {"error": str(e)}

def evaluate_url(url: str):
    print(f"[*] Scanning standards for {url} ...")
    # Headers
    headers_data = run_node("headers_scan.js", url, timeout=25)
    # Axe (WCAG)
    axe_data = run_node("axe_scan.js", url, timeout=40)
    # Lighthouse (CWV/SEO)
    lh_data = run_node("lighthouse_scan.js", url, timeout=60)
    # Wappalyzer
    wapp_data = run_node("wappalyzer_scan.js", url, timeout=30)
    
    rep = build_standards_report(
        url=url,
        axe_data=axe_data,
        lighthouse_data=lh_data,
        headers_data=headers_data,
        wappalyzer_data=wapp_data,
    )
    return rep

def main():
    print(f"Loading {EXCEL_PATH} ...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["Standards Score Report"]
    
    # Map URLs to row indices
    url_to_row = {}
    for r in range(3, 603):
        u = ws.cell(r, 2).value
        if u:
            url_to_row[u.strip()] = r
            
    # Process DEAD_SITES
    font_arial = Font(name="Arial", size=10)
    font_bold = Font(name="Arial", size=10, bold=True)
    border_thin = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0")
    )
    fill_dead = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    font_dead = Font(name="Arial", size=10, bold=True, color="842029")
    
    for dead_u, reason in DEAD_SITES.items():
        if dead_u in url_to_row:
            r = url_to_row[dead_u]
            print(f"[!] Updating DEAD site: Row {r} -> {dead_u} ({reason})")
            ws.cell(r, 4, value=f"=F{r}/71")
            ws.cell(r, 5, value=f'=IF(F{r}=0,"เว็บปิด",IF(D{r}>=0.9,"ดีมาก",IF(D{r}>=0.7,"ดี",IF(D{r}>=0.5,"พอใช้","ต้องปรับปรุง"))))')
            ws.cell(r, 6, value=0)
            ws.cell(r, 7, value=0)
            ws.cell(r, 8, value=71)
            
            # 4 Standards details
            # WCAG (cols 9-12)
            ws.cell(r, 9, value=f"=J{r}/37")
            ws.cell(r, 10, value=0)
            ws.cell(r, 11, value=0)
            ws.cell(r, 12, value=37)
            # CWV (cols 13-16)
            ws.cell(r, 13, value=f"=N{r}/9")
            ws.cell(r, 14, value=0)
            ws.cell(r, 15, value=0)
            ws.cell(r, 16, value=9)
            # NCSA (cols 17-20)
            ws.cell(r, 17, value=f"=R{r}/11")
            ws.cell(r, 18, value=0)
            ws.cell(r, 19, value=0)
            ws.cell(r, 20, value=11)
            # OWASP (cols 21-24)
            ws.cell(r, 21, value=f"=V{r}/14")
            ws.cell(r, 22, value=0)
            ws.cell(r, 23, value=0)
            ws.cell(r, 24, value=14)
            
            # Apply dead styles across all cols
            for c in range(1, 25):
                cell = ws.cell(r, c)
                cell.fill = fill_dead
                if c in [1, 2, 4, 5]:
                    cell.font = font_dead
                    
    # Rescan live sites
    for u in RESCAN_SITES:
        if u not in url_to_row:
            print(f"Skipping {u} (not in excel)")
            continue
        r = url_to_row[u]
        rep = evaluate_url(u)
        
        ov = rep["summary"]
        ov["score_pct"] = round(ov["passed"] / 71 * 100)
        stds = {s["id"]: s for s in rep["standards"]}
        wcag = stds.get("wcag", {})
        cwv = stds.get("cwv", {})
        ncsa = stds.get("ncsa", {})
        owasp = stds.get("owasp", {})
        
        print(f"  -> Result: Passed={ov['passed']}/71 ({ov['score_pct']}%), WCAG={wcag.get('passed',0)}/37, CWV={cwv.get('passed',0)}/9, NCSA={ncsa.get('passed',0)}/11, OWASP={owasp.get('passed',0)}/14")
        
        ws.cell(r, 4, value=f"=F{r}/71")
        ws.cell(r, 5, value=f'=IF(F{r}=0,"เว็บปิด",IF(D{r}>=0.9,"ดีมาก",IF(D{r}>=0.7,"ดี",IF(D{r}>=0.5,"พอใช้","ต้องปรับปรุง"))))')
        ws.cell(r, 6, value=ov["passed"])
        ws.cell(r, 7, value=ov["warning"])
        ws.cell(r, 8, value=ov["failed"])
        
        # WCAG
        ws.cell(r, 9, value=f"=J{r}/37")
        ws.cell(r, 10, value=wcag.get("passed", 0))
        ws.cell(r, 11, value=wcag.get("warning", 0))
        ws.cell(r, 12, value=wcag.get("failed", 0))
        
        # CWV
        ws.cell(r, 13, value=f"=N{r}/9")
        ws.cell(r, 14, value=cwv.get("passed", 0))
        ws.cell(r, 15, value=cwv.get("warning", 0))
        ws.cell(r, 16, value=cwv.get("failed", 0))
        
        # NCSA
        ws.cell(r, 17, value=f"=R{r}/11")
        ws.cell(r, 18, value=ncsa.get("passed", 0))
        ws.cell(r, 19, value=ncsa.get("warning", 0))
        ws.cell(r, 20, value=ncsa.get("failed", 0))
        
        # OWASP
        ws.cell(r, 21, value=f"=V{r}/14")
        ws.cell(r, 22, value=owasp.get("passed", 0))
        ws.cell(r, 23, value=owasp.get("warning", 0))
        ws.cell(r, 24, value=owasp.get("failed", 0))
        
        # Determine grade and color
        pct = ov["score_pct"]
        if ov["passed"] == 0:
            fill_c = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
            font_c = Font(name="Arial", size=10, bold=True, color="842029")
        elif pct >= 90:
            fill_c = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
            font_c = Font(name="Arial", size=10, bold=True, color="155724")
        elif pct >= 70:
            fill_c = PatternFill(start_color="CFE2FF", end_color="CFE2FF", fill_type="solid")
            font_c = Font(name="Arial", size=10, bold=True, color="084298")
        elif pct >= 50:
            fill_c = PatternFill(start_color="CFF4FC", end_color="CFF4FC", fill_type="solid")
            font_c = Font(name="Arial", size=10, bold=True, color="055160")
        else:
            fill_c = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
            font_c = Font(name="Arial", size=10, bold=True, color="664D03")
            
        for c in range(1, 25):
            cell = ws.cell(r, c)
            cell.fill = fill_c
            if c in [4, 5]:
                cell.font = font_c
                
    wb.save(EXCEL_PATH)
    print("Successfully updated standards_score_results.xlsx!")

if __name__ == "__main__":
    main()
