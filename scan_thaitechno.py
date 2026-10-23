"""
scan_thaitechno.py
Runs the complete Web Project standards assessment on thaitechno.net (Row 203)
and updates standards_score_results.xlsx with the exact compliance scores and grade.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
BASE_DIR = Path(__file__).resolve().parent
BACKEND_DIR = BASE_DIR / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from repair_user_targets import evaluate_url

URL = "http://thaitechno.net"
ROW = 203
EXCEL_PATH = BASE_DIR / "standards_score_results.xlsx"

print(f"[*] Starting full Web Project standards evaluation for {URL} ...")
rep = evaluate_url(URL)
summary = rep["summary"]
stds = {s["id"]: s for s in rep["standards"]}

wcag = stds.get("wcag", {})
cwv = stds.get("cwv", {})
ncsa = stds.get("ncsa", {})
owasp = stds.get("owasp", {})

passed = summary["passed"]
warning = summary["warning"]
failed = summary["failed"]
score_pct = round(passed / 71 * 100)

print(f"[+] Assessment Complete!")
print(f"    Total Passed: {passed}/71 ({score_pct}%)")
print(f"    WCAG 2.1: Passed={wcag.get('passed',0)}/37, Warning={wcag.get('warning',0)}, Failed={wcag.get('failed',0)}")
print(f"    CWV & SEO: Passed={cwv.get('passed',0)}/9, Warning={cwv.get('warning',0)}, Failed={cwv.get('failed',0)}")
print(f"    สกมช. NCSA: Passed={ncsa.get('passed',0)}/11, Warning={ncsa.get('warning',0)}, Failed={ncsa.get('failed',0)}")
print(f"    OWASP: Passed={owasp.get('passed',0)}/14, Warning={owasp.get('warning',0)}, Failed={owasp.get('failed',0)}")

# Determine grade and color
# >= 90: ดีมาก (เขียว)
# >= 70: ดี (น้ำเงิน)
# >= 50: พอใช้ (ฟ้า)
# < 50: ต้องปรับปรุง (เหลือง)
if passed == 0:
    grade_label = "เว็บปิด"
    fill_hex = "F8D7DA"
    font_hex = "842029"
elif score_pct >= 90:
    grade_label = "ดีมาก"
    fill_hex = "D4EDDA"
    font_hex = "155724"
elif score_pct >= 70:
    grade_label = "ดี"
    fill_hex = "CFE2FF"
    font_hex = "084298"
elif score_pct >= 50:
    grade_label = "พอใช้"
    fill_hex = "CFF4FC"
    font_hex = "055160"
else:
    grade_label = "ต้องปรับปรุง"
    fill_hex = "FFF3CD"
    font_hex = "664D03"

print(f"    Grade: {grade_label} (Color: {fill_hex})")

# Update Excel
print(f"[*] Updating {EXCEL_PATH} Row {ROW} ...")
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb["Standards Score Report"]

# Verify row URL
row_url = ws.cell(ROW, 2).value
print(f"    Row {ROW} URL in Excel: {row_url}")

ws.cell(ROW, 4, value=f"=F{ROW}/71")
ws.cell(ROW, 5, value=f'=IF(F{ROW}=0,"เว็บปิด",IF(D{ROW}>=0.9,"ดีมาก",IF(D{ROW}>=0.7,"ดี",IF(D{ROW}>=0.5,"พอใช้","ต้องปรับปรุง"))))')
ws.cell(ROW, 6, value=passed)
ws.cell(ROW, 7, value=warning)
ws.cell(ROW, 8, value=failed)

# WCAG
ws.cell(ROW, 9, value=f"=J{ROW}/37")
ws.cell(ROW, 10, value=wcag.get("passed", 0))
ws.cell(ROW, 11, value=wcag.get("warning", 0))
ws.cell(ROW, 12, value=wcag.get("failed", 0))

# CWV
ws.cell(ROW, 13, value=f"=N{ROW}/9")
ws.cell(ROW, 14, value=cwv.get("passed", 0))
ws.cell(ROW, 15, value=cwv.get("warning", 0))
ws.cell(ROW, 16, value=cwv.get("failed", 0))

# NCSA
ws.cell(ROW, 17, value=f"=R{ROW}/11")
ws.cell(ROW, 18, value=ncsa.get("passed", 0))
ws.cell(ROW, 19, value=ncsa.get("warning", 0))
ws.cell(ROW, 20, value=ncsa.get("failed", 0))

# OWASP
ws.cell(ROW, 21, value=f"=V{ROW}/14")
ws.cell(ROW, 22, value=owasp.get("passed", 0))
ws.cell(ROW, 23, value=owasp.get("warning", 0))
ws.cell(ROW, 24, value=owasp.get("failed", 0))

fill_cell = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
font_grade = Font(name="Arial", size=10, bold=True, color=font_hex)

for c in range(1, 25):
    ws.cell(ROW, c).fill = fill_cell
    if c in [4, 5]:
        ws.cell(ROW, c).font = font_grade

wb.save(EXCEL_PATH)
print(f"[✓] Successfully updated thaitechno.net in {EXCEL_PATH}!")
