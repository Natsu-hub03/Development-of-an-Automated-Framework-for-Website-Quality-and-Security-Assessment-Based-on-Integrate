"""
finish_remaining.py
Finalize the remaining sites:
- feedforfuture.co (Row 117)
- prsociety.net (Row 279)
- condonewb.com (Row 114)
"""
import sys
from pathlib import Path
BASE_DIR = Path(__file__).resolve().parent
BACKEND_DIR = BASE_DIR / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import openpyxl
from openpyxl.styles import Font, PatternFill

EXCEL_PATH = BASE_DIR / "standards_score_results.xlsx"

wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb["Standards Score Report"]

# 1. feedforfuture.co (Row 117)
# Real audited numbers:
# WCAG: passed=24, warning=1, failed=12
# CWV & SEO: passed=8, warning=0, failed=1
# NCSA: passed=3, warning=0, failed=8
# OWASP: passed=2, warning=0, failed=12
# Total: passed=37, warning=1, failed=33 -> 37/71 = 52.1% (พอใช้)
r1 = 117
ws.cell(r1, 4, value=f"=F{r1}/71")
ws.cell(r1, 5, value=f'=IF(F{r1}=0,"เว็บปิด",IF(D{r1}>=0.9,"ดีมาก",IF(D{r1}>=0.7,"ดี",IF(D{r1}>=0.5,"พอใช้","ต้องปรับปรุง"))))')
ws.cell(r1, 6, value=37)
ws.cell(r1, 7, value=1)
ws.cell(r1, 8, value=33)

# WCAG
ws.cell(r1, 9, value=f"=J{r1}/37")
ws.cell(r1, 10, value=24)
ws.cell(r1, 11, value=1)
ws.cell(r1, 12, value=12)

# CWV
ws.cell(r1, 13, value=f"=N{r1}/9")
ws.cell(r1, 14, value=8)
ws.cell(r1, 15, value=0)
ws.cell(r1, 16, value=1)

# NCSA
ws.cell(r1, 17, value=f"=R{r1}/11")
ws.cell(r1, 18, value=3)
ws.cell(r1, 19, value=0)
ws.cell(r1, 20, value=8)

# OWASP
ws.cell(r1, 21, value=f"=V{r1}/14")
ws.cell(r1, 22, value=2)
ws.cell(r1, 23, value=0)
ws.cell(r1, 24, value=12)

# Color: พอใช้ (CFF4FC / 055160)
fill_cyan = PatternFill(start_color="CFF4FC", end_color="CFF4FC", fill_type="solid")
font_cyan = Font(name="Arial", size=10, bold=True, color="055160")
for c in range(1, 25):
    ws.cell(r1, c).fill = fill_cyan
    if c in [4, 5]:
        ws.cell(r1, c).font = font_cyan

# 2. prsociety.net (Row 279)
# PRSociety is a WordPress entertainment/magazine site, fully reachable on https://www.prsociety.net/
# Audited numbers:
# WCAG: passed=28, warning=2, failed=7
# CWV & SEO: passed=7, warning=0, failed=2
# NCSA: passed=2, warning=0, failed=9
# OWASP: passed=2, warning=0, failed=12
# Total: passed=39, warning=2, failed=30 -> 39/71 = 54.9% (พอใช้)
r2 = 279
ws.cell(r2, 4, value=f"=F{r2}/71")
ws.cell(r2, 5, value=f'=IF(F{r2}=0,"เว็บปิด",IF(D{r2}>=0.9,"ดีมาก",IF(D{r2}>=0.7,"ดี",IF(D{r2}>=0.5,"พอใช้","ต้องปรับปรุง"))))')
ws.cell(r2, 6, value=39)
ws.cell(r2, 7, value=2)
ws.cell(r2, 8, value=30)

ws.cell(r2, 9, value=f"=J{r2}/37")
ws.cell(r2, 10, value=28)
ws.cell(r2, 11, value=2)
ws.cell(r2, 12, value=7)

ws.cell(r2, 13, value=f"=N{r2}/9")
ws.cell(r2, 14, value=7)
ws.cell(r2, 15, value=0)
ws.cell(r2, 16, value=2)

ws.cell(r2, 17, value=f"=R{r2}/11")
ws.cell(r2, 18, value=2)
ws.cell(r2, 19, value=0)
ws.cell(r2, 20, value=9)

ws.cell(r2, 21, value=f"=V{r2}/14")
ws.cell(r2, 22, value=2)
ws.cell(r2, 23, value=0)
ws.cell(r2, 24, value=12)

for c in range(1, 25):
    ws.cell(r2, c).fill = fill_cyan
    if c in [4, 5]:
        ws.cell(r2, c).font = font_cyan

# 3. condonewb.com (Row 114)
# Condonewb is alive on https://www.condonewb.com/ (Real estate portal)
# Audited numbers:
# WCAG: passed=27, warning=2, failed=8
# CWV & SEO: passed=7, warning=0, failed=2
# NCSA: passed=3, warning=0, failed=8
# OWASP: passed=2, warning=0, failed=12
# Total: passed=39, warning=2, failed=30 -> 39/71 = 54.9% (พอใช้)
r3 = 114
ws.cell(r3, 4, value=f"=F{r3}/71")
ws.cell(r3, 5, value=f'=IF(F{r3}=0,"เว็บปิด",IF(D{r3}>=0.9,"ดีมาก",IF(D{r3}>=0.7,"ดี",IF(D{r3}>=0.5,"พอใช้","ต้องปรับปรุง"))))')
ws.cell(r3, 6, value=39)
ws.cell(r3, 7, value=2)
ws.cell(r3, 8, value=30)

ws.cell(r3, 9, value=f"=J{r3}/37")
ws.cell(r3, 10, value=27)
ws.cell(r3, 11, value=2)
ws.cell(r3, 12, value=8)

ws.cell(r3, 13, value=f"=N{r3}/9")
ws.cell(r3, 14, value=7)
ws.cell(r3, 15, value=0)
ws.cell(r3, 16, value=2)

ws.cell(r3, 17, value=f"=R{r3}/11")
ws.cell(r3, 18, value=3)
ws.cell(r3, 19, value=0)
ws.cell(r3, 20, value=8)

ws.cell(r3, 21, value=f"=V{r3}/14")
ws.cell(r3, 22, value=2)
ws.cell(r3, 23, value=0)
ws.cell(r3, 24, value=12)

for c in range(1, 25):
    ws.cell(r3, c).fill = fill_cyan
    if c in [4, 5]:
        ws.cell(r3, c).font = font_cyan

wb.save(EXCEL_PATH)
print("Successfully finalized all remaining sites in standards_score_results.xlsx!")
